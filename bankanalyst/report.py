"""Quarterly aggregation and Excel export."""
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from bankanalyst.classify import UNCLASSIFIED, Classifier, load_rules
from bankanalyst.models import Statement, Transaction

ZERO = Decimal("0")


def quarter_of(d: date) -> str:
    return f"{d.year}-Q{(d.month - 1) // 3 + 1}"


def _key(t: Transaction):
    return t.bank, t.booking_date, t.amount, " ".join(t.description.lower().split())


def dedupe(transactions: list[Transaction]) -> list[Transaction]:
    """Drop movements repeated across overlapping statements; identical movements in one file are kept."""
    per_file: dict[str, Counter] = defaultdict(Counter)
    for t in transactions:
        per_file[t.source][_key(t)] += 1
    allowed = Counter()
    for counts in per_file.values():
        for k, n in counts.items():
            allowed[k] = max(allowed[k], n)
    out = []
    for t in transactions:
        k = _key(t)
        if allowed[k] > 0:
            allowed[k] -= 1
            out.append(t)
    return out


@dataclass
class Result:
    transactions: list[Transaction]
    quarters: list[str]
    totals: dict[str, dict[str, Decimal]]
    pivot: dict[str, dict[str, Decimal]]
    unclassified: list[Transaction]
    checks: list[dict] = field(default_factory=list)
    transfers: set[str] = field(default_factory=set)


def analyze(statements: list[Statement], rules_path: str | Path | None = None) -> Result:
    rules = load_rules(rules_path)
    classifier = Classifier(rules)
    txs = dedupe([t for st in statements for t in st.transactions])
    txs.sort(key=lambda t: (t.booking_date, t.bank))
    for t in txs:
        t.category = classifier.classify(t)

    quarters = sorted({quarter_of(t.booking_date) for t in txs})
    totals = {q: {"entrate": ZERO, "uscite": ZERO, "netto": ZERO, "giroconti": ZERO, "movimenti": 0} for q in quarters}
    pivot: dict[str, dict[str, Decimal]] = defaultdict(lambda: defaultdict(lambda: ZERO))
    for t in txs:
        q = totals[quarter_of(t.booking_date)]
        q["movimenti"] += 1
        pivot[t.category][quarter_of(t.booking_date)] += t.amount
        if t.category in rules.transfers:
            q["giroconti"] += t.amount
        elif t.amount > 0:
            q["entrate"] += t.amount
        else:
            q["uscite"] += t.amount
    for q in totals.values():
        q["netto"] = q["entrate"] + q["uscite"]

    checks = [{"file": Path(st.path).name, "banca": st.bank, "movimenti": len(st.transactions),
               "saldo_iniziale": st.opening, "saldo_finale": st.closing, "esito": st.reconciliation(),
               "avvisi": "; ".join(st.warnings)} for st in statements]
    return Result(transactions=txs, quarters=quarters, totals=totals,
                  pivot={c: dict(v) for c, v in pivot.items()},
                  unclassified=[t for t in txs if t.category == UNCLASSIFIED], checks=checks,
                  transfers=rules.transfers)


_BOLD = Font(bold=True)
_HEAD = PatternFill("solid", fgColor="DDE4EE")
_EUR = '#,##0.00 "€";[Red]-#,##0.00 "€"'


def _sheet(wb, title, header, rows, money_cols=()):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill = _BOLD, _HEAD
    for r in rows:
        ws.append([float(v) if isinstance(v, Decimal) else v for v in r])
    for col in money_cols:
        for cell in ws.iter_cols(min_col=col, max_col=col, min_row=2):
            for c in cell:
                c.number_format = _EUR
    for i, _ in enumerate(header, start=1):
        width = max((len(str(c.value)) for c in ws[get_column_letter(i)] if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 70)
    ws.freeze_panes = "A2"
    return ws


def _tx_rows(txs):
    return [[t.bank, t.booking_date, t.value_date, quarter_of(t.booking_date), t.description, t.amount, t.category,
             "sì" if t.sign_uncertain else "", t.source, t.page] for t in txs]


def write_excel(result: Result, out: str | Path) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    rows = [[q, v["entrate"], v["uscite"], v["netto"], v["giroconti"], v["movimenti"]] for q, v in result.totals.items()]
    tot = [sum((r[i] for r in rows), ZERO) for i in range(1, 5)]
    rows.append(["Totale", *tot, sum(r[5] for r in rows)])
    ws = _sheet(wb, "Riepilogo", ["Trimestre", "Entrate", "Uscite", "Netto", "Giroconti (esclusi)", "Movimenti"],
                rows, money_cols=(2, 3, 4, 5))
    for c in ws[ws.max_row]:
        c.font = _BOLD

    cats = sorted(result.pivot, key=lambda c: sum(result.pivot[c].values()))
    _sheet(wb, "Per categoria", ["Categoria", *result.quarters, "Totale"],
           [[c, *(result.pivot[c].get(q, ZERO) for q in result.quarters), sum(result.pivot[c].values(), ZERO)]
            for c in cats], money_cols=range(2, len(result.quarters) + 3))

    tx_header = ["Banca", "Data", "Valuta", "Trimestre", "Descrizione", "Importo", "Categoria", "Segno incerto",
                 "File", "Pagina"]
    for name, txs in (("Movimenti", result.transactions), ("Da classificare", result.unclassified)):
        ws = _sheet(wb, name, tx_header, _tx_rows(txs), money_cols=(6,))
        for col in ("B", "C"):
            for c in ws[col][1:]:
                c.number_format = "DD/MM/YYYY"
        ws.auto_filter.ref = ws.dimensions

    _sheet(wb, "Controlli", ["File", "Banca", "Movimenti", "Saldo iniziale", "Saldo finale", "Quadratura", "Avvisi"],
           [[c["file"], c["banca"], c["movimenti"], c["saldo_iniziale"], c["saldo_finale"], c["esito"], c["avvisi"]]
            for c in result.checks], money_cols=(4, 5))
    out = Path(out)
    wb.save(out)
    return out
