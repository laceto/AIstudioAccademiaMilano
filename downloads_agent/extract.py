"""Read the start of a document as plain text. Never raises: returns (text, note)."""
from __future__ import annotations

import logging
import warnings
from pathlib import Path
from typing import Tuple

logging.getLogger("pypdf").setLevel(logging.CRITICAL)   # damaged PDFs are reported via the note
warnings.filterwarnings("ignore", module="openpyxl")

MAX_CHARS = 3000          # enough for a small local model; keeps prompts short
MAX_FILE_BYTES = 50 * 1024 * 1024
MAX_SHEET_ROWS = 40
READABLE = {".pdf", ".docx", ".txt", ".md", ".csv", ".xls", ".xlsx"}


def _plain(path: Path) -> str:
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        return f.read(MAX_CHARS)


def _docx(path: Path) -> str:
    import docx
    out, size = [], 0
    for p in docx.Document(str(path)).paragraphs:
        out.append(p.text)
        size += len(p.text)
        if size >= MAX_CHARS:
            break
    return "\n".join(out)


def _xlsx(path: Path) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    try:
        rows = []
        for row in wb.worksheets[0].iter_rows(max_row=MAX_SHEET_ROWS, values_only=True):
            rows.append(" | ".join("" if c is None else str(c) for c in row))
        return "\n".join(rows)
    finally:
        wb.close()


def _xls(path: Path) -> str:
    import xlrd
    sheet = xlrd.open_workbook(str(path), on_demand=True).sheet_by_index(0)
    rows = []
    for r in range(min(sheet.nrows, MAX_SHEET_ROWS)):
        rows.append(" | ".join(str(c) for c in sheet.row_values(r)))
    return "\n".join(rows)


def _pdf(path: Path) -> str:
    from pypdf import PdfReader
    out, size = [], 0
    for page in PdfReader(str(path)).pages:
        t = page.extract_text() or ""
        out.append(t)
        size += len(t)
        if size >= MAX_CHARS:
            break
    return "\n".join(out)


_READERS = {".txt": _plain, ".md": _plain, ".csv": _plain, ".docx": _docx,
            ".xlsx": _xlsx, ".xls": _xls, ".pdf": _pdf}


def extract_text(path) -> Tuple[str, str]:
    path = Path(path)
    try:
        if path.stat().st_size > MAX_FILE_BYTES:
            return "", "too large to read"
        text = _READERS[path.suffix.lower()](path).strip()
    except Exception as e:  # corrupt, encrypted, wrong format ...
        return "", f"unreadable: {type(e).__name__}"
    if not text:
        return "", "no extractable text"
    return text[:MAX_CHARS], ""
